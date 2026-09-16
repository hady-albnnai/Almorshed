#!/usr/bin/env python3
"""تحويل نوطة الأستاذ (Markdown محدود) إلى ملف Word عربي RTL.

الاستعمال:
    python3 tools/note_to_docx.py content/authoring/notes/FADAA-NOTE-U1.md out.docx

يدعم: عناوين # ## ###، جداول |، اقتباسات >، قوائم - و1.، خط عريض **، كود `…`
(يُعرض بخط مونو LTR)، وفواصل ---. الاتجاه RTL على مستوى الفقرة والجدول،
والمعادلات داخل ` ` تُوسم LTR حتى لا تنقلب الرموز.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ARABIC_FONT = "Arial"
MONO_FONT = "Consolas"
GOLD = RGBColor(0x8A, 0x6D, 0x1F)
GREY = RGBColor(0x55, 0x55, 0x55)


# ───────────────────────── مساعدات XML للاتجاه ─────────────────────────

def _set_rtl(paragraph, rtl: bool = True) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")


def _set_run_rtl(run, rtl: bool) -> None:
    rPr = run._r.get_or_add_rPr()
    el = rPr.find(qn("w:rtl"))
    if el is None:
        el = OxmlElement("w:rtl")
        rPr.append(el)
    el.set(qn("w:val"), "1" if rtl else "0")


def _set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)
    if size:
        run.font.size = Pt(size)
        # حجم الخط للنص المركب (العربي) أيضاً
        szCs = rPr.find(qn("w:szCs"))
        if szCs is None:
            szCs = OxmlElement("w:szCs")
            rPr.append(szCs)
        szCs.set(qn("w:val"), str(int(size * 2)))


def _table_rtl(table) -> None:
    tblPr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)


def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


# ───────────────────────── تنسيق النص المضمّن ─────────────────────────

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, size: float = 11, base_bold: bool = False,
               color: RGBColor | None = None) -> None:
    """يضيف نصاً مع دعم **عريض** و`كود`؛ الكود LTR بخط مونو."""
    parts = INLINE_RE.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            _set_run_font(run, ARABIC_FONT, size)
            _set_run_rtl(run, True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, MONO_FONT, max(size - 0.5, 9))
            _set_run_rtl(run, False)
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            run.bold = base_bold
        else:
            run = paragraph.add_run(part)
            run.bold = base_bold
            _set_run_font(run, ARABIC_FONT, size)
            _set_run_rtl(run, True)
        if color is not None and not (part.startswith("`")):
            run.font.color.rgb = color


def add_paragraph(doc, text: str, size: float = 11, bold: bool = False,
                  align=WD_ALIGN_PARAGRAPH.RIGHT, indent_cm: float = 0.0,
                  color: RGBColor | None = None, space_after: float = 4):
    p = doc.add_paragraph()
    p.alignment = align
    _set_rtl(p, True)
    if indent_cm:
        p.paragraph_format.right_indent = Cm(indent_cm)
    p.paragraph_format.space_after = Pt(space_after)
    add_inline(p, text, size=size, base_bold=bold, color=color)
    return p


def add_heading(doc, text: str, level: int) -> None:
    sizes = {1: 20, 2: 16, 3: 13.5, 4: 12}
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl(p, True)
    p.paragraph_format.space_before = Pt(14 if level <= 2 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    clean = re.sub(r"\*\*", "", text)
    add_inline(p, clean, size=sizes.get(level, 12), base_bold=True,
               color=GOLD if level <= 2 else None)
    if level == 1:
        _bottom_border(p)


def _bottom_border(p) -> None:
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "8A6D1F")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_quote(doc, text: str) -> None:
    """اقتباس > … : صندوق مظلل (تنبيه السلم / ملاحظة)."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_rtl(table)
    cell = table.rows[0].cells[0]
    warn = text.startswith("⚠️") or "تنبيه" in text[:20]
    _shade(cell, "FFF4D6" if warn else "EEF3F8")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl(p, True)
    add_inline(p, text, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, rows: list[list[str]]) -> None:
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _table_rtl(table)
    for i, row in enumerate(rows):
        for j in range(ncols):
            txt = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl(p, True)
            add_inline(p, txt, size=10, base_bold=(i == 0))
            if i == 0:
                _shade(cell, "E8E0C8")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ───────────────────────── المحلّل ─────────────────────────

def parse_table_block(lines: list[str]) -> list[list[str]]:
    rows = []
    for ln in lines:
        if re.match(r"^\|\s*-{2,}", ln):
            continue  # سطر الفواصل
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def convert(md_path: Path, out_path: Path) -> None:
    doc = Document()
    # هوامش وخط افتراضي
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = ARABIC_FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:cs"), ARABIC_FONT)

    lines = md_path.read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s == "---":
            p = doc.add_paragraph()
            _bottom_border(p)
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            add_heading(doc, m.group(2), len(m.group(1)))
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, parse_table_block(block))
            continue
        if s.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip()[1:].strip())
                i += 1
            add_quote(doc, " ".join(b for b in block if b))
            continue
        m = re.match(r"^[-•]\s+(.*)$", s)
        if m:
            add_paragraph(doc, "• " + m.group(1), indent_cm=0.6, space_after=2)
            i += 1
            continue
        m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if m:
            add_paragraph(doc, f"{m.group(1)}. {m.group(2)}", indent_cm=0.6, space_after=2)
            i += 1
            continue
        add_paragraph(doc, s)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
    print("→", sys.argv[2])
